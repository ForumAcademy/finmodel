"""
Собирает тестовый кейс tests/cases/derbenevskaya_legacy.yaml из исходного Excel:
  - project_inputs: значения параметров проекта (как в исходнике)
  - capex_legacy: вбитые суммы/ставки по статьям и ручные ряды распределения (помесячно не пересчитываются — квартальные)
  - sales_legacy: темпы продаж по кварталам
  - legacy_outputs: итоговые значения исходника (для сверки, НЕ эталон — многие неверны, см. docs/02_legacy_audit.md)
"""
import re
from datetime import datetime
from pathlib import Path
import datetime as dt
import openpyxl, yaml
from openpyxl.utils import get_column_letter as gl

ROOT = Path(__file__).resolve().parents[1]
X = ROOT / "legacy" / "Кальк Саевой привязка КОД.xlsx"
wv = openpyxl.load_workbook(X, data_only=True)
t, b, cf, ps, es, d = (wv[n] for n in ["ТЭПы", "Бюджет ", "CF1 ", "План продаж", "Эскроу", "Dashboard"])
wf = openpyxl.load_workbook(X)  # формулы: какие ячейки вбиты числом

def num(x):
    if isinstance(x, float):
        return round(x, 6)
    return x

def series(ws, row, c1, c2):
    return [num(ws.cell(row, c).value or 0) for c in range(c1, c2 + 1)]

quarters = [cf.cell(2, c).value for c in range(6, 46)]  # F..AS
quarters = [q.date().isoformat() if isinstance(q, dt.datetime) else None for q in quarters]

capex = yaml.safe_load(open(ROOT / "data" / "capex_items.yaml"))["items"]
capex_legacy = []
for it in capex:
    lg = it["legacy"]
    br = lg.get("budget_row")
    entry = {"item_id": it["item_id"], "budget_row": br}
    if br:
        entry.update({"rate_D": num(b.cell(br, 4).value), "qty_E": num(b.cell(br, 5).value), "amount_F": num(b.cell(br, 6).value)})
    sched_rows = [r for r in lg.get("cf_rows", []) if isinstance(cf.cell(r, 1).value, str) and "темп" in cf.cell(r, 1).value.lower()]
    if sched_rows:
        entry["manual_schedule_quarterly"] = {"cf_row": sched_rows[0], "values_F_to_AS": series(cf, sched_rows[0], 6, 45),
                                              "sum": num(sum(series(cf, sched_rows[0], 6, 45)))}
    elif (it["base"] == "формула" or it["schedule_rule"] == "follow_sales") and lg.get("cf_rows"):
        # статья со своей формулой или от выручки без ряда «темп»: суммы CF1 по кварталам (распределение исходника для расчёта «как в исходном Excel»)
        r = lg["cf_rows"][0]
        entry["cf_amounts_quarterly"] = {"cf_row": r, "values_F_to_AS": series(cf, r, 6, 45), "sum": num(sum(series(cf, r, 6, 45)))}
    capex_legacy.append(entry)

QUARTER_END = {1: "03-31", 2: "06-30", 3: "09-30", 4: "12-31"}


def quarter_end(text):
    """«4 кв 2025» → «2025-12-31»: последний день квартала (решение владельца продукта 25.09.2026)."""
    m = re.search(r"(\d)\s*кв\s*(\d{4})\s*$", str(text).strip())
    return f"{m.group(2)}-{QUARTER_END[int(m.group(1))]}"


def legacy_milestones(t):
    """
    Вехи исходника (ТЭПы!C6:C12) заданы текстом по проекту в целом → даты по очередям.
    Приобретение участка, РНС и старт продаж — общие для всех очередей; окончание СМР (одна дата
    на проект) — у последней очереди; даты РНВ — по порядку очередей. Прочие вехи в исходнике не заданы.
    """
    phases = int(t["C17"].value)
    rnv = [datetime.strptime(x.strip(), "%d.%m.%Y").date().isoformat() for x in str(t["C10"].value).split(";") if x.strip()]
    rows = []
    for p in range(1, phases + 1):
        rows.append({
            "phase": p,
            "land_acquired": quarter_end(t["C6"].value),
            "rns_date": quarter_end(t["C7"].value),
            "sales_start": quarter_end(t["C8"].value),
            "construction_end": quarter_end(t["C9"].value) if p == phases else None,
            "rnv_date": rnv[p - 1] if p <= len(rnv) else None,
        })
    return rows


def legacy_milestones_note(t):
    rnv = [x.strip() for x in str(t["C10"].value).split(";") if x.strip()]
    phases = int(t["C17"].value)
    extra = rnv[phases:]
    return (f"В исходнике {len(rnv)} даты РНВ на {phases} очереди; по очередям взяты первые {phases}"
            + (f", не распределены: {', '.join(extra)}" if extra else "") + ". Окончание СМР — одна дата на проект.")


case = {
    "case_id": "derbenevskaya_legacy",
    "description": "Входные данные исходного файла «Кальк Саевой привязка КОД.xlsx» (Дербеневская наб., Москва). Используется для сверки и регрессионных тестов.",
    "source_file": "legacy/Кальк Саевой привязка КОД.xlsx",
    "project_inputs": {
        "GEN.PROJECT_NAME": t["C2"].value,
        "GEN.REGION_CODE": "77",
        "GEN.HOUSING_CLASS": "бизнес",
        "GEN.MODEL_START_DATE": "2025-12-31",
        "GEN.PHASES_COUNT": 3,
        "GEN.PROJECT_STAGE": "концепция",
        "LAND.AREA": t["C15"].value,
        "LAND.CADASTRAL_VALUE": t["C13"].value,
        "TEP.GFA_ABOVE": t["C19"].value,
        "TEP.RES_GFA": num(t["C20"].value),
        "TEP.NONRES_GFA": t["C21"].value,
        "TEP.APT_AREA": t["C22"].value,
        "TEP.COMM_AREA": num(t["C23"].value),
        "TEP.MOP_AREA": num(t["C26"].value),
        "TEP.GFA_BELOW": num(t["E49"].value),
        "TEP.PARKING_COUNT_OVERRIDE": t["C27"].value,
        "TEP.PARKING_AREA_PER_SPACE": t["D49"].value,
        "TEP.LANDSCAPE_SHARE": t["D29"].value,
        "TEP.ROAD_SHARE": t["D30"].value,
        "TEP.GREEN_SHARE": t["D31"].value,
        "TEP.APT_MIX": [{"type_name": t.cell(r, 2).value.strip(), "count": t.cell(r, 6).value, "avg_area": t.cell(r, 4).value,
                          "start_price": t.cell(r, 7).value, "parking_norm": t.cell(r, 9).value} for r in (41, 42, 43)],
        "SALES.PRODUCTS.ПСН": {"stock_area": num(t["C23"].value), "avg_lot": t["D48"].value, "start_price": t["G48"].value},
        "SALES.PRODUCTS.машино-места": {"price_per_m2": t["G49"].value, "area_per_space": t["D49"].value,
                                         "price_per_space_calc": num(t["G49"].value * t["D49"].value)},
        "TIME.MILESTONES_TEXT": {k: t[c].value for k, c in [("land", "C6"), ("rns", "C7"), ("sales_start", "C8"),
                                  ("smr_end", "C9"), ("rnv", "C10"), ("sales_end", "C11"), ("pf_end", "C12")]},
        "TIME.MILESTONES": legacy_milestones(t),
        "TIME.MILESTONES_NOTE": legacy_milestones_note(t),
        "SALES.PAYMENT_MIX": {"installment": es["C5"].value, "mortgage": es["C6"].value, "full": es["C7"].value, "down_payment": es["C8"].value,
                               "installment_quarters": es["D3"].value},
        "OPEX.MARKETING_RATE": b["D51"].value, "OPEX.BROKERAGE_RATE": b["D52"].value, "OPEX.DEV_FEE_RATE": b["D48"].value,
        "LAND.AGENT_FEE_RATE": b["D19"].value,
        "FIN.EQUITY_SHARE": cf["D132"].value, "FIN.RATE_PREFERENTIAL": cf["D116"].value,
        "FIN.RATE_BASE_legacy": cf["D117"].value, "FIN.KEY_RATE_legacy": cf["D115"].value,
        "FIN.RATE_DISCOUNT_COEF": cf["D118"].value, "FIN.RATE_MIN": cf["D119"].value,
        "FIN.ESCROW_RESERVE_RATE": cf["D92"].value, "FIN.FEE_ARRANGEMENT": cf["D111"].value, "FIN.LEGACY_LIMIT": num(b["F67"].value),
        "FIN.COLLATERAL_DISCOUNT": cf["D102"].value, "VAL.COST_OF_EQUITY_legacy": cf["D139"].value,
        "TAX.LAND_RATE_legacy": cf["D84"].value,
    },
    "timeline_quarters_F_to_AS": quarters,
    "sales_legacy": {
        "pace_units_quarterly_from_1q2026": {
            "Тип 1": series(ps, 28, 5, 40), "Тип 2": series(ps, 33, 5, 40), "Тип 3": series(ps, 38, 5, 40),
            "ПСН": series(ps, 43, 5, 40), "Машино-места": series(ps, 48, 5, 40)},
        "price_growth_quarterly": ps["E30"].value,
    },
    "capex_legacy": capex_legacy,
    # Ячейки исходника для проверок расчёта «как в исходном Excel» (расхождения внутри Excel → предупреждения)
    "legacy_checks": {
        "tep": {"apt_area_C22": t["C22"].value, "psn_stock_C23": t["C23"].value, "saleable_area_C35": t["C35"].value,
                "saleable_area_C35_formula": wf["ТЭПы"]["C35"].value if str(wf["ТЭПы"]["C35"].value).startswith("=") else None,
                "psn_stock_C48": t["C48"].value, "psn_lot_D48": t["D48"].value, "escrow_release_C12": t["C12"].value,
                "gfa_above_C19": t["C19"].value, "res_gfa_C20": num(t["C20"].value), "res_gfa_C20_formula": wf["ТЭПы"]["C20"].value,
                "nonres_gfa_C21": t["C21"].value},
        "sales_plan": {"revenue_row25_from_1q2026": series(ps, 25, 5, 40), "psn_lots_row43_sum": num(ps["D43"].value)},
        "budget": {"marketing_rate_D51": b["D51"].value, "marketing_F51": num(b["F51"].value), "marketing_F51_formula": wf["Бюджет "]["F51"].value,
                   "brokerage_rate_D52": b["D52"].value, "brokerage_F52": num(b["F52"].value),
                   "contingency_D42": num(b["D42"].value), "contingency_E42": b["E42"].value, "contingency_F42": num(b["F42"].value),
                   "contingency_F42_formula": wf["Бюджет "]["F42"].value},
        "cf1": {"revenue_row15": series(cf, 15, 6, 45), "revenue_row15_last_formula_col": max(c for c in range(6, 46) if wf["CF1 "].cell(15, c).value) - 6,
                "brokerage_row78": series(cf, 78, 6, 45), "marketing_row79": series(cf, 79, 6, 45),
                "escrow_release_row6": series(cf, 6, 6, 45), "escrow_release_row6_typed": [c - 6 for c in range(6, 46) if not str(wf["CF1 "].cell(6, c).value).startswith("=")],
                "escrow_deposit_row8": series(cf, 8, 6, 45), "escrow_date_D6_formula": wf["CF1 "]["D6"].value},
        # Кредит CF1 (строки 98–132, по кварталам F…AS; знаки — как в CF1) и лимит Бюджет!F67
        "fin": {"limit_F67": num(b["F67"].value), "limit_F67_formula": wf["Бюджет "]["F67"].value,
                "draw_row98": series(cf, 98, 6, 45), "repaid_escrow_row100": series(cf, 100, 6, 45), "debt_row106": series(cf, 106, 6, 45),
                "interest_row107": series(cf, 107, 6, 45), "pik_paid_row108": series(cf, 108, 6, 45), "accrued_row110": series(cf, 110, 6, 45),
                "fee_row111": series(cf, 111, 6, 45), "flow_row113": series(cf, 113, 6, 45), "k1_row122": series(cf, 122, 6, 45),
                "rate_row125": series(cf, 125, 6, 45), "equity_row132": series(cf, 132, 6, 45), "eff_rate_D128": cf["D128"].value},
    },
    "legacy_outputs_for_reference": {
        "_warning": "Значения исходника. Многие рассчитаны с ошибками (docs/02_legacy_audit.md) — использовать только для сверки, не как эталон.",
        "revenue_sales_plan": num(ps["D25"].value),
        "revenue_cf1": num(cf["C15"].value),
        "budget_total_costs": num(b["F59"].value),
        "budget_finrez": num(b["F61"].value),
        "budget_smr": num(b["F30"].value),
        "cf1_costs": num(cf["C18"].value),
        "cf1_interest": num(cf["C107"].value),
        "cf1_draws": num(cf["C98"].value),
        "cf1_fcfe": num(cf["C138"].value),
        "dashboard_irr_project": num(d["C44"].value),
        "dashboard_irr_equity": num(d["C45"].value),
        "dashboard_ltv_max": num(d["C46"].value),
        "dashboard_credit_term_quarters": num(d["C38"].value),
        "cost_per_m2_budget": num(b["F69"].value),
        "markup_budget": num(b["F71"].value),
    },
}

case["reconciliation_targets"] = {
    "_how": "Ядро в расчёте «как в исходном Excel» (квартальный ручной темп из sales_legacy, рост цены 2%/кв, цены из ТЭП) обязано воспроизвести эти числа исходника с точностью 1 руб / 0,01 м². Эти значения исходника считаются арифметически верными.",
    "F.TEP.APT_TYPE_AREA_sum": round(t["E44"].value, 3),
    "F.TEP.PARKING_REQUIRED_by_legacy_norm": 2785,
    "F.TEP.LANDSCAPE_AREA": round(t["C29"].value, 3),
    "F.TEP.GFA_BELOW_EST_legacy_parking": round(t["E49"].value, 3),
    "sales_value_type1": round(ps["D27"].value, 2),
    "sales_value_type2": round(ps["D32"].value, 2),
    "sales_value_type3": round(ps["D37"].value, 2),
    "sales_value_psn": round(ps["D42"].value, 2),
    "sales_value_parking": round(ps["D47"].value, 2),
    "sales_value_total": round(ps["D25"].value, 2),
}
case["expected_differences_after_fix"] = {
    "_how": "После исправлений новая модель ОБЯЗАНА отличаться от исходника в этих местах (иначе ошибка не устранена).",
    "revenue_in_cf_equals_sales_plan": "CF выручка = 118 130 288 024,50 (исходник CF1 — 117 514 091 710,14)",
    "uds_in_cf": "статья ROADS_UDS 535 620 851 присутствует в CF (исходник — нет)",
    "brokerage_in_cf": "брокеридж в CF = 3,5% × выручка по месяцам продаж (исходник CF1 — 2 733 176 548,21 вместо 4 134 560 080,86)",
    "other_smr_in_cf": "OTHER_SMR 16 904 952,57 распределены (исходник — 0)",
    "contingency": "резерв = ставка × база, а не E42 + D42",
    "psn_sold_le_stock": "продано ПСН ≤ 10 322 м² (исходник 10 888,6 м²)",
    "credit_term_nonzero": "срок кредита > 0 (исходник Dashboard!C38 = 0)",
}

out = ROOT / "tests" / "cases" / "derbenevskaya_legacy.yaml"
out.parent.mkdir(parents=True, exist_ok=True)
yaml.safe_dump(case, open(out, "w", encoding="utf-8"), allow_unicode=True, sort_keys=False, width=200)
print("written", out)
